from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from hashlib import sha256
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from modules.weight_progress.models import WeightMeasurement


SUPPORTED_SHEET_NAME = "weight"
UPLOAD_MAX_BYTES = 8 * 1024 * 1024
OVERLAP_WINDOW = timedelta(minutes=5)
OVERLAP_WEIGHT_TOLERANCE_LB = 0.2


def _normalize_header(value):
    return str(value or "").strip().lower()


def _is_blank(value):
    return value is None or str(value).strip() == ""


def _lb_to_kg(value):
    return round(float(value) / 2.2046226218, 6)


def _format_date_label(value):
    if not value:
        return "--"
    return value.strftime("%b %d, %Y").replace(" 0", " ")


def _format_datetime_label(value):
    if not value:
        return "--"
    return value.strftime("%b %d, %Y at %I:%M %p").replace(" 0", " ").replace(" 0", " ")


@dataclass
class ParsedWorkbookRow:
    row_number: int
    source_timestamp: datetime
    source_timestamp_text: str
    weight_lb: float
    fat_mass_lb: float | None = None
    bone_mass_lb: float | None = None
    muscle_mass_lb: float | None = None
    hydration_lb: float | None = None
    comments: str | None = None
    body_fat_percent: float | None = None
    exact_hash: str = ""
    issues: list[str] = field(default_factory=list)

    def to_measurement(self, *, import_source, source_label, source_filename):
        body_fat = self.body_fat_percent
        if body_fat is None and self.fat_mass_lb is not None and self.weight_lb > 0:
            body_fat = round((self.fat_mass_lb / self.weight_lb) * 100, 4)

        metadata = {
            "import_source": import_source,
            "source_label": source_label,
            "source_filename": source_filename,
            "sheet": SUPPORTED_SHEET_NAME,
            "row_number": self.row_number,
            "source_timestamp": self.source_timestamp_text,
            "source_units": "lb",
        }

        return WeightMeasurement.create(
            captured_at=datetime.now().isoformat(sep=" ", timespec="microseconds"),
            source_timestamp=self.source_timestamp.isoformat(sep=" ", timespec="microseconds"),
            source_entity="sensor.withings_weight",
            weight_kg=_lb_to_kg(self.weight_lb),
            body_fat_percent=body_fat,
            fat_mass_kg=_lb_to_kg(self.fat_mass_lb) if self.fat_mass_lb is not None else None,
            fat_free_mass_kg=None,
            muscle_mass_kg=_lb_to_kg(self.muscle_mass_lb) if self.muscle_mass_lb is not None else None,
            bone_mass_kg=_lb_to_kg(self.bone_mass_lb) if self.bone_mass_lb is not None else None,
            hydration_kg=_lb_to_kg(self.hydration_lb) if self.hydration_lb is not None else None,
            comments=self.comments,
            import_source=import_source,
            source_label=source_label,
            imported_at=datetime.now().isoformat(sep=" ", timespec="microseconds"),
            reading_hash=self.exact_hash,
            metadata=metadata,
        )


@dataclass
class WorkbookImportResult:
    filename: str
    worksheet: str
    rows_found: int
    valid_measurements: int
    invalid_rows: int
    exact_duplicates: int
    likely_overlaps: int
    measurements_before_journey_start: int
    measurements_on_or_after_journey_start: int
    earliest_measurement: str
    latest_measurement: str
    proposed_journey_baseline: str
    proposed_starting_weight: str
    proposed_journey_date: str
    journey_start_date: str | None
    journey_start_label: str
    journey_start_source: str
    journey_start_value_lb: float | None
    journey_start_value_kg: float | None
    current_journey_count: int
    all_history_count: int
    errors: list[dict] = field(default_factory=list)
    measurements: list[ParsedWorkbookRow] = field(default_factory=list)

    def to_preview_dict(self):
        return {
            "filename": self.filename,
            "worksheet": self.worksheet,
            "rows_found": self.rows_found,
            "valid_measurements": self.valid_measurements,
            "invalid_rows": self.invalid_rows,
            "exact_duplicates": self.exact_duplicates,
            "likely_overlaps": self.likely_overlaps,
            "measurements_before_journey_start": self.measurements_before_journey_start,
            "measurements_on_or_after_journey_start": self.measurements_on_or_after_journey_start,
            "earliest_measurement": self.earliest_measurement,
            "latest_measurement": self.latest_measurement,
            "proposed_journey_baseline": self.proposed_journey_baseline,
            "proposed_starting_weight": self.proposed_starting_weight,
            "proposed_journey_date": self.proposed_journey_date,
            "journey_start_date": self.journey_start_date,
            "journey_start_label": self.journey_start_label,
            "journey_start_source": self.journey_start_source,
            "current_journey_count": self.current_journey_count,
            "all_history_count": self.all_history_count,
            "errors": self.errors,
        }


class WithingsWorkbookImporter:
    HEADER_MAP = {
        "date": "timestamp",
        "weight (lb)": "weight_lb",
        "fat mass (lb)": "fat_mass_lb",
        "bone mass (lb)": "bone_mass_lb",
        "muscle mass (lb)": "muscle_mass_lb",
        "hydration (lb)": "hydration_lb",
        "comments": "comments",
        "body fat (%)": "body_fat_percent",
        "body fat percent": "body_fat_percent",
    }

    REQUIRED_HEADERS = {"timestamp", "weight_lb"}

    def __init__(self, log):
        self.log = log

    def preview(self, path, existing_measurements, journey_start_date, source_label, filename):
        parsed_rows, errors, worksheet_name = self._parse_rows(path)
        return self._build_result(
            parsed_rows=parsed_rows,
            errors=errors,
            worksheet_name=worksheet_name,
            existing_measurements=existing_measurements,
            journey_start_date=journey_start_date,
            source_label=source_label,
            filename=filename,
        )

    def build_measurements(self, path, existing_measurements, journey_start_date, source_label, filename):
        result = self.preview(path, existing_measurements, journey_start_date, source_label, filename)
        measurements = [
            row.to_measurement(
                import_source="withings_xlsx",
                source_label=source_label,
                source_filename=filename,
            )
            for row in result.measurements
        ]
        return result, measurements

    def _build_result(self, parsed_rows, errors, worksheet_name, existing_measurements, journey_start_date, source_label, filename):
        journey_date = self._parse_journey_date(journey_start_date)
        existing_rows = list(existing_measurements or [])
        exact_seen = set()
        valid_rows = []
        exact_duplicates = 0
        likely_overlaps = 0
        before_journey = 0
        on_or_after_journey = 0
        ordered_rows = []
        existing_hashes = {measurement.reading_hash for measurement in existing_rows if measurement.reading_hash}

        for row in parsed_rows:
            if row.issues:
                continue

            if row.exact_hash in exact_seen or row.exact_hash in existing_hashes:
                exact_duplicates += 1
                continue
            exact_seen.add(row.exact_hash)

            if self._is_likely_overlap(row, existing_rows):
                likely_overlaps += 1
                continue

            valid_rows.append(row)
            ordered_rows.append(row)

        ordered_rows.sort(key=lambda item: item.source_timestamp)
        for row in ordered_rows:
            if journey_date and row.source_timestamp.date() < journey_date:
                before_journey += 1
            else:
                on_or_after_journey += 1

        proposed_baseline = self._select_baseline(ordered_rows, journey_date)
        proposed_weight = proposed_baseline.weight_lb if proposed_baseline else None
        proposed_date = proposed_baseline.source_timestamp.date() if proposed_baseline else journey_date

        result = WorkbookImportResult(
            filename=filename,
            worksheet=worksheet_name or SUPPORTED_SHEET_NAME,
            rows_found=len(parsed_rows),
            valid_measurements=len(valid_rows),
            invalid_rows=len(errors),
            exact_duplicates=exact_duplicates,
            likely_overlaps=likely_overlaps,
            measurements_before_journey_start=before_journey,
            measurements_on_or_after_journey_start=on_or_after_journey,
            earliest_measurement=_format_datetime_label(ordered_rows[0].source_timestamp) if ordered_rows else "--",
            latest_measurement=_format_datetime_label(ordered_rows[-1].source_timestamp) if ordered_rows else "--",
            proposed_journey_baseline=_format_datetime_label(proposed_baseline.source_timestamp) if proposed_baseline else "--",
            proposed_starting_weight=f"{proposed_weight:.1f} lb" if proposed_weight is not None else "--",
            proposed_journey_date=_format_date_label(proposed_date) if proposed_date else "--",
            journey_start_date=journey_date.isoformat() if journey_date else None,
            journey_start_label=_format_date_label(journey_date) if journey_date else "Since HomePulse tracking began",
            journey_start_source=f"Starting Weight — {_format_date_label(journey_date)}" if journey_date else "Since HomePulse tracking began",
            journey_start_value_lb=proposed_weight,
            journey_start_value_kg=_lb_to_kg(proposed_weight) if proposed_weight is not None else None,
            current_journey_count=on_or_after_journey if journey_date else len(valid_rows),
            all_history_count=len(valid_rows),
            errors=errors,
            measurements=ordered_rows,
        )
        return result

    def _parse_rows(self, path):
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        try:
            worksheet = self._get_worksheet(workbook)
            headers, header_row_number = self._read_headers(worksheet)
            parsed_rows = []
            errors = []
            for row_number, row in self._iter_data_rows(worksheet, header_row_number + 1):
                parsed_row = self._parse_row(row_number, row, headers, workbook)
                if parsed_row is None:
                    continue
                if parsed_row.issues:
                    errors.append({"row": row_number, "messages": list(parsed_row.issues)})
                parsed_rows.append(parsed_row)
            return parsed_rows, errors, worksheet.title
        finally:
            workbook.close()

    def _get_worksheet(self, workbook):
        for sheet in workbook.worksheets:
            if _normalize_header(sheet.title) == SUPPORTED_SHEET_NAME:
                return sheet
        raise ValueError("Workbook must include a worksheet named 'weight'.")

    def _read_headers(self, worksheet):
        for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=10), start=1):
            values = [cell.value for cell in row]
            if all(_is_blank(value) for value in values):
                continue
            headers = {}
            for index, value in enumerate(values):
                mapped = self.HEADER_MAP.get(_normalize_header(value))
                if mapped:
                    headers[mapped] = index
            missing = [header for header in self.REQUIRED_HEADERS if header not in headers]
            if missing:
                raise ValueError("Workbook is missing required headers: Date and Weight (lb).")
            return headers, row_number
        raise ValueError("Workbook does not contain a header row.")

    def _iter_data_rows(self, worksheet, start_row):
        for row_number, row in enumerate(worksheet.iter_rows(min_row=start_row), start=start_row):
            if all(_is_blank(cell.value) for cell in row):
                continue
            yield row_number, row

    def _parse_row(self, row_number, row, headers, workbook):
        issues = []
        timestamp = self._parse_timestamp(row[headers["timestamp"]].value if "timestamp" in headers else None, workbook)
        if timestamp is None:
            issues.append("Date is missing or invalid.")

        weight_lb = self._parse_number(row[headers["weight_lb"]].value if "weight_lb" in headers else None)
        if weight_lb is None or weight_lb <= 0:
            issues.append("Weight (lb) is missing or invalid.")

        fat_mass_lb = self._parse_optional_number(row, headers, "fat_mass_lb", issues)
        bone_mass_lb = self._parse_optional_number(row, headers, "bone_mass_lb", issues)
        muscle_mass_lb = self._parse_optional_number(row, headers, "muscle_mass_lb", issues)
        hydration_lb = self._parse_optional_number(row, headers, "hydration_lb", issues)
        body_fat_percent = self._parse_optional_number(row, headers, "body_fat_percent", issues)
        comments = self._parse_text(row[headers["comments"]].value) if "comments" in headers else None

        if issues:
            return ParsedWorkbookRow(
                row_number=row_number,
                source_timestamp=timestamp or datetime.min,
                source_timestamp_text="",
                weight_lb=weight_lb or 0,
                fat_mass_lb=fat_mass_lb,
                bone_mass_lb=bone_mass_lb,
                muscle_mass_lb=muscle_mass_lb,
                hydration_lb=hydration_lb,
                comments=comments,
                body_fat_percent=body_fat_percent,
                exact_hash="",
                issues=issues,
            )

        timestamp_text = timestamp.isoformat(sep=" ", timespec="microseconds")
        exact_hash = self._build_exact_hash(
            timestamp_text,
            weight_lb,
            body_fat_percent,
            fat_mass_lb,
            bone_mass_lb,
            muscle_mass_lb,
            hydration_lb,
            comments,
        )
        if body_fat_percent is None and fat_mass_lb is not None and weight_lb > 0:
            body_fat_percent = round((fat_mass_lb / weight_lb) * 100, 4)

        return ParsedWorkbookRow(
            row_number=row_number,
            source_timestamp=timestamp,
            source_timestamp_text=timestamp_text,
            weight_lb=weight_lb,
            fat_mass_lb=fat_mass_lb,
            bone_mass_lb=bone_mass_lb,
            muscle_mass_lb=muscle_mass_lb,
            hydration_lb=hydration_lb,
            comments=comments,
            body_fat_percent=body_fat_percent,
            exact_hash=exact_hash,
        )

    @staticmethod
    def _parse_optional_number(row, headers, key, issues):
        if key not in headers:
            return None
        value = row[headers[key]].value
        if _is_blank(value):
            return None
        number = WithingsWorkbookImporter._parse_number(value)
        if number is None:
            issues.append(f"{key.replace('_', ' ').title()} is invalid.")
        return number

    @staticmethod
    def _parse_text(value):
        if _is_blank(value):
            return None
        return str(value).strip()

    @staticmethod
    def _parse_number(value):
        if _is_blank(value):
            return None
        try:
            return float(str(value).strip().replace(",", ""))
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _parse_timestamp(value, workbook):
        if _is_blank(value):
            return None
        if isinstance(value, datetime):
            return value.replace(tzinfo=None)
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        if isinstance(value, (int, float)):
            try:
                return from_excel(value, workbook.epoch)
            except Exception:
                return None
        text = str(value).strip().replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(text)
            return parsed.astimezone().replace(tzinfo=None) if parsed.tzinfo else parsed
        except ValueError:
            return None

    @staticmethod
    def _build_exact_hash(timestamp_text, weight_lb, body_fat_percent, fat_mass_lb, bone_mass_lb, muscle_mass_lb, hydration_lb, comments):
        payload = {
            "timestamp": timestamp_text,
            "weight_lb": round(weight_lb, 4),
            "body_fat_percent": round(body_fat_percent, 4) if body_fat_percent is not None else None,
            "fat_mass_lb": round(fat_mass_lb, 4) if fat_mass_lb is not None else None,
            "bone_mass_lb": round(bone_mass_lb, 4) if bone_mass_lb is not None else None,
            "muscle_mass_lb": round(muscle_mass_lb, 4) if muscle_mass_lb is not None else None,
            "hydration_lb": round(hydration_lb, 4) if hydration_lb is not None else None,
            "comments": comments or "",
        }
        return sha256(str(sorted(payload.items())).encode("utf-8")).hexdigest()

    @staticmethod
    def _parse_journey_date(value):
        if not value:
            return None
        if isinstance(value, date):
            return value
        text = str(value).strip()
        if not text:
            return None
        try:
            return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            return None

    def _is_likely_overlap(self, row, existing_rows):
        row_ts = row.source_timestamp
        row_weight_lb = row.weight_lb
        for measurement in existing_rows:
            if measurement.source_timestamp is None or measurement.weight_kg is None:
                continue
            existing_ts = self._parse_existing_timestamp(measurement.source_timestamp)
            if existing_ts is None:
                continue
            if abs((row_ts - existing_ts).total_seconds()) > OVERLAP_WINDOW.total_seconds():
                continue
            existing_weight_lb = measurement.weight_kg * 2.2046226218
            if abs(existing_weight_lb - row_weight_lb) <= OVERLAP_WEIGHT_TOLERANCE_LB:
                return True
        return False

    @staticmethod
    def _parse_existing_timestamp(value):
        if not value:
            return None
        text = str(value).strip().replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(text)
            return parsed.astimezone().replace(tzinfo=None) if parsed.tzinfo else parsed
        except ValueError:
            return None

    @staticmethod
    def _select_baseline(rows, journey_date):
        if not rows:
            return None
        if journey_date is None:
            return rows[0]
        for row in rows:
            if row.source_timestamp.date() >= journey_date:
                return row
        return rows[0]
